"""
Document Builder Utility
Role: Generate professional, ATS-friendly DOCX and PDF files.
"""

from typing import Dict, Any, List
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from reportlab.lib.pagesizes import LETTER
from reportlab.lib.units import inch
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, ListFlowable, ListItem

class DocumentBuilder:
    """
    Handles creation and formatting of MS Word documents.
    """

    def __init__(self):
        self.doc = Document()
        self._setup_styles()

    def _setup_styles(self):
        """Configure document styles for ATS readability"""
        # Set margins (standard 1 inch)
        for section in self.doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Standard font
        style = self.doc.styles['Normal']
        font = style.font
        font.name = 'Calibri'
        font.size = Pt(11)

    def create_cv(self, cv_data: Dict[str, Any], output_path: str):
        """
        Generate a CV document from structured data.

        Args:
            cv_data: Dictionary containing 'personal_info', 'experience', 'education', 'skills'
            output_path: File path to save the DOCX
        """
        try:
            # 1. Header (Name & Contact)
            self._add_header(cv_data.get('personal_info', {}))

            # 2. Professional Summary
            if 'summary' in cv_data:
                self._add_section_title("PROFESSIONAL SUMMARY")
                self.doc.add_paragraph(cv_data['summary'])

            # 3. Skills
            if 'skills' in cv_data:
                self._add_section_title("CORE SKILLS")
                self._add_skills(cv_data['skills'])

            # 4. Experience
            if 'experience' in cv_data:
                self._add_section_title("PROFESSIONAL EXPERIENCE")
                for role in cv_data['experience']:
                    self._add_experience_item(role)

            # 5. Education
            if 'education' in cv_data:
                self._add_section_title("EDUCATION")
                for edu in cv_data['education']:
                    self._add_education_item(edu)
            
            # Save
            self.doc.save(output_path)
            print(f"✅ Document saved to: {output_path}")

        except Exception as e:
            print(f"❌ Failed to create document: {e}")
            raise

    def _add_header(self, info: Dict[str, str]):
        """Add personal info header"""
        name = info.get('name', 'Candidate Name')
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        run = p.add_run(name)
        run.bold = True
        run.font.size = Pt(20)
        run.font.color.rgb = RGBColor(0, 0, 0) # Black

        # Contact line
        contact_parts = []
        if info.get('email'): contact_parts.append(info['email'])
        if info.get('phone'): contact_parts.append(info['phone'])
        if info.get('linkedin'): contact_parts.append(info['linkedin'])
        if info.get('location'): contact_parts.append(info['location'])
        
        if contact_parts:
            p = self.doc.add_paragraph(" | ".join(contact_parts))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.runs[0].font.size = Pt(10)

    def _add_section_title(self, title: str):
        """Add a standardized section header"""
        p = self.doc.add_paragraph()
        p.space_before = Pt(12)
        p.space_after = Pt(6)
        
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(12)
        run.font.name = 'Calibri'
        
        # Add bottom border style (hacky in python-docx, usually simple underline is safer for ATS)
        # Using simple underline for safety
        run.underline = True

    def _add_skills(self, skills: Any):
        """Format skills section"""
        if isinstance(skills, list):
            # If simple list, join with bullets or pipes? 
            # ATS prefers comma separated or bullet points.
            self.doc.add_paragraph(", ".join(skills))
        elif isinstance(skills, dict):
            # Categorized skills
            for category, items in skills.items():
                p = self.doc.add_paragraph()
                run = p.add_run(f"{category}: ")
                run.bold = True
                p.add_run(", ".join(items))

    def _add_experience_item(self, role: Dict[str, Any]):
        """Add a job role"""
        # Title line
        p = self.doc.add_paragraph()
        p.space_before = Pt(8)
        
        # Company Name (Bold)
        company = role.get('company', '')
        r1 = p.add_run(company)
        r1.bold = True
        
        # Location (Right align if possible, but keep simple for ATS)
        location = role.get('location', '')
        if location:
            p.add_run(f" — {location}")

        # Job Title & Dates
        p2 = self.doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(2)
        title = role.get('title', '')
        dates = role.get('dates', '')
        
        r2 = p2.add_run(title)
        r2.italic = True
        if dates:
            p2.add_run(f" | {dates}")

        # Bullets
        achievements = role.get('achievements', role.get('responsibilities', []))
        for item in achievements:
            self.doc.add_paragraph(item, style='List Bullet')

    def _add_education_item(self, edu: Dict[str, Any]):
        """Add education item"""
        p = self.doc.add_paragraph()
        p.space_before = Pt(6)
        
        school = edu.get('school', '')
        degree = edu.get('degree', '')
        dates = edu.get('dates', '')
        
        r = p.add_run(school)
        r.bold = True
        p.add_run(f" — {degree}")
        if dates:
            p.add_run(f" ({dates})")

    def create_cover_letter(self, letter_body: str, profile: Dict[str, Any], output_path: str):
        """
        Generate a Cover Letter document.
        
        Args:
            letter_body: The text content of the letter
            profile: Candidate profile (for header)
            output_path: File path to save
        """
        try:
            # Re-initialize doc for new file
            self.doc = Document()
            self._setup_styles()
            
            # 1. Header (Same as CV)
            self._add_header(profile.get('personal_info', {}))
            
            # 2. Spacing
            self.doc.add_paragraph().space_after = Pt(24)
            
            # 3. Body
            # Split by newlines to create proper paragraphs
            for paragraph in letter_body.split('\n'):
                if paragraph.strip():
                    p = self.doc.add_paragraph(paragraph.strip())
                    p.paragraph_format.space_after = Pt(12)
            
            # Save
            self.doc.save(output_path)
            print(f"✅ Cover Letter saved to: {output_path}")
            
        except Exception as e:
            print(f"❌ Failed to create cover letter: {e}")
            raise

    # ------------------------------------------------------------------
    # PDF generation (reportlab) — mirrors the DOCX layout above so users
    # can choose either format without the content diverging.
    # ------------------------------------------------------------------

    def _pdf_styles(self):
        styles = getSampleStyleSheet()
        styles.add(ParagraphStyle(
            name='NameHeader', parent=styles['Title'], fontSize=20,
            alignment=TA_CENTER, spaceAfter=4,
        ))
        styles.add(ParagraphStyle(
            name='ContactLine', parent=styles['Normal'], fontSize=10,
            alignment=TA_CENTER, spaceAfter=12,
        ))
        styles.add(ParagraphStyle(
            name='SectionTitle', parent=styles['Heading2'], fontSize=12,
            spaceBefore=12, spaceAfter=6, underline=True,
        ))
        return styles

    def _pdf_header_flowables(self, info: Dict[str, str], styles) -> List[Any]:
        name = info.get('name', 'Candidate Name')
        contact_parts = [
            info[key] for key in ('email', 'phone', 'linkedin', 'location')
            if info.get(key)
        ]
        flowables = [Paragraph(name, styles['NameHeader'])]
        if contact_parts:
            flowables.append(Paragraph(" | ".join(contact_parts), styles['ContactLine']))
        return flowables

    def create_cv_pdf(self, cv_data: Dict[str, Any], output_path: str):
        """
        Generate a CV as a PDF, mirroring create_cv()'s DOCX layout.

        Args:
            cv_data: Dictionary containing 'personal_info', 'experience', 'education', 'skills'
            output_path: File path to save the PDF
        """
        try:
            styles = self._pdf_styles()
            story: List[Any] = []
            story.extend(self._pdf_header_flowables(cv_data.get('personal_info', {}), styles))

            if 'summary' in cv_data:
                story.append(Paragraph("PROFESSIONAL SUMMARY", styles['SectionTitle']))
                story.append(Paragraph(cv_data['summary'], styles['Normal']))

            skills = cv_data.get('skills')
            if skills:
                story.append(Paragraph("CORE SKILLS", styles['SectionTitle']))
                if isinstance(skills, dict):
                    for category, items in skills.items():
                        story.append(Paragraph(f"<b>{category}:</b> {', '.join(items)}", styles['Normal']))
                elif isinstance(skills, list):
                    story.append(Paragraph(", ".join(skills), styles['Normal']))

            experience = cv_data.get('experience')
            if experience:
                story.append(Paragraph("PROFESSIONAL EXPERIENCE", styles['SectionTitle']))
                for role in experience:
                    company = role.get('company', '')
                    location = role.get('location', '')
                    header = f"<b>{company}</b>" + (f" — {location}" if location else "")
                    story.append(Paragraph(header, styles['Normal']))

                    title = role.get('title', '')
                    dates = role.get('dates', '')
                    subheader = f"<i>{title}</i>" + (f" | {dates}" if dates else "")
                    story.append(Paragraph(subheader, styles['Normal']))

                    achievements = role.get('achievements', role.get('responsibilities', []))
                    if achievements:
                        story.append(ListFlowable(
                            [ListItem(Paragraph(item, styles['Normal'])) for item in achievements],
                            bulletType='bullet',
                        ))
                    story.append(Spacer(1, 8))

            education = cv_data.get('education')
            if education:
                story.append(Paragraph("EDUCATION", styles['SectionTitle']))
                for edu in education:
                    school = edu.get('school', '')
                    degree = edu.get('degree', '')
                    dates = edu.get('dates', '')
                    line = f"<b>{school}</b> — {degree}" + (f" ({dates})" if dates else "")
                    story.append(Paragraph(line, styles['Normal']))

            doc = SimpleDocTemplate(
                output_path, pagesize=LETTER,
                topMargin=inch, bottomMargin=inch, leftMargin=inch, rightMargin=inch,
            )
            doc.build(story)
            print(f"✅ PDF saved to: {output_path}")

        except Exception as e:
            print(f"❌ Failed to create PDF: {e}")
            raise

    def create_cover_letter_pdf(self, letter_body: str, profile: Dict[str, Any], output_path: str):
        """
        Generate a Cover Letter as a PDF, mirroring create_cover_letter()'s DOCX layout.

        Args:
            letter_body: The text content of the letter
            profile: Candidate profile (for header)
            output_path: File path to save the PDF
        """
        try:
            styles = self._pdf_styles()
            story: List[Any] = []
            story.extend(self._pdf_header_flowables(profile.get('personal_info', {}), styles))
            story.append(Spacer(1, 24))

            for paragraph in letter_body.split('\n'):
                if paragraph.strip():
                    story.append(Paragraph(paragraph.strip(), styles['Normal']))
                    story.append(Spacer(1, 12))

            doc = SimpleDocTemplate(
                output_path, pagesize=LETTER,
                topMargin=inch, bottomMargin=inch, leftMargin=inch, rightMargin=inch,
            )
            doc.build(story)
            print(f"✅ Cover Letter PDF saved to: {output_path}")

        except Exception as e:
            print(f"❌ Failed to create cover letter PDF: {e}")
            raise
